"""Journey: opening a file attached to a chat shows what it holds, before and after sending.

A file attached in the composer, and the same file on the sent message, opens a dialog with its
name, its size and the text the server read out of it. Its Preview tab shows a CSV as a table of
its rows, each sheet of a spreadsheet on a tab of its own, a Word document as its rendered page and
Markdown with its formatting. A PowerPoint slide draws its title where the file aligns it: left,
right or, when the file leaves it open, centred. Its pie and doughnut charts draw their slices, and
a horizontal bar chart lays its bars across while a column chart keeps them upright. A PDF shows its
extracted text; the Preview tab's PDF viewer says "Failed to load PDF." in the suite's Chromium 131
for any PDF, so it is not driven here. An attached picture opens full size from the sent message and
downloads as the bytes that were uploaded. Switching a long document to Using Entire Document in its
dialog sends the model the whole text where focused retrieval sends a few pieces of it, and a file's
switch shows its own mode after another attachment is removed.

Discriminates: passes on the dev ebc6add67 build. In a frontend build whose file dialog shows
"No content" for the text, no longer takes a CSV for a spreadsheet, hides the sheet tabs, never
loads a Word document, shows Markdown as plain text and ignores the Using Entire Document switch,
and whose picture preview never opens, every test fails. In the pre-fix build of 0.12.4
(a944564c2 reverted) the remaining file shows the removed file's Using Entire Document. The
PowerPoint title test fails on the dev 75358502c build (#32873: the left-aligned title starts 42%
across the slide) and passes with the fix 8fa77c828 (#32876). The PowerPoint chart test fails
on the dev 52f51ea4f build (#32884: the pie and doughnut slides stay blank and the longest
horizontal bar spans 13% of the slide) and passes with the fix f0f0e4ab1, which is not merged yet.
"""

from __future__ import annotations

import io
import json
import uuid
from pathlib import Path

import pytest
from PIL import Image
from playwright.sync_api import Locator, Page, expect

from harness import upstream as reply
from integration.retrieval.test_document_loaders import (
    DOCX,
    PPTX,
    XLSX,
    docx_bytes,
    pdf_bytes,
)
from utils.cached_chat import attach
from utils.chat_ui import conversation, expect_reply, send

pytestmark = [pytest.mark.journey, pytest.mark.requires_browser, pytest.mark.requires_source]

PROCESSING_TIMEOUT_MS = 30_000


def file_chip(scope: Locator, name: str) -> Locator:
    return scope.get_by_role("button").filter(has_text=name).first


def open_file(page: Page, chip: Locator, name: str) -> Locator:
    """Click a file's chip and return its dialog once the server's reading shows."""
    chip.click()
    dialog = page.get_by_role("dialog").filter(has_text=name)
    expect(dialog.get_by_text("Extracted Lines")).to_be_visible(timeout=PROCESSING_TIMEOUT_MS)
    return dialog


def attach_and_open(page: Page, name: str, content: str | bytes, mime_type: str) -> Locator:
    attach(page, name, content, mime_type)
    chip = file_chip(page.locator("form"), name)
    expect(chip).to_be_visible()
    # the chip spins until the server has read the file; its dialog loads the text only once
    expect(chip.locator(".spinner_ajPY")).to_have_count(0, timeout=PROCESSING_TIMEOUT_MS)
    return open_file(page, chip, name)


def show_preview(dialog: Locator) -> None:
    dialog.get_by_role("button", name="Preview", exact=True).click()


def ask_with_attachments(page: Page, upstream, question: str) -> None:
    upstream.queue(reply.text("Noted.", match=reply.answering(question)))
    send(page, question)
    expect_reply(page, "Noted.")


def test_a_text_file_opens_with_its_size_and_text_in_the_composer_and_the_message(
    page_for, make_user, upstream
):
    page = page_for(make_user())
    text = "The harbour master is Ingrid.\nThe ferry leaves at seven."

    dialog = attach_and_open(page, "harbour.txt", text, "text/plain")
    expect(dialog.get_by_text(f"{len(text.encode())}.0 B")).to_be_visible()
    expect(dialog.get_by_text("2 Extracted Lines")).to_be_visible()
    expect(dialog.get_by_text("The ferry leaves at seven.")).to_be_visible()
    page.keyboard.press("Escape")
    expect(dialog).to_have_count(0)

    ask_with_attachments(page, upstream, f"who runs the harbour? {uuid.uuid4().hex[:6]}")
    sent = open_file(page, file_chip(conversation(page), "harbour.txt"), "harbour.txt")
    expect(sent.get_by_text("The harbour master is Ingrid.")).to_be_visible()


def test_a_csv_previews_as_a_table_of_its_rows(page_for, make_user):
    page = page_for(make_user())
    rows = "route,departs\nNorth pier,07:15\nSouth pier,08:40\n"

    dialog = attach_and_open(page, "ferries.csv", rows, "text/csv")
    show_preview(dialog)

    expect(dialog.get_by_text("3 Rows")).to_be_visible()
    expect(dialog.get_by_role("row").filter(has_text="South pier")).to_contain_text("08:40")
    expect(dialog.get_by_role("row").filter(has_text="North pier")).to_contain_text("07:15")


def test_each_sheet_of_a_spreadsheet_previews_on_its_own_tab(page_for, make_user):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    workbook.active.title = "Tides"
    workbook.active.append(["Monday", "high at noon"])
    workbook.create_sheet("Ferries").append(["North pier", "07:15"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    page = page_for(make_user())

    dialog = attach_and_open(page, "harbour.xlsx", buffer.getvalue(), XLSX)
    show_preview(dialog)

    expect(dialog.get_by_role("row").filter(has_text="Monday")).to_contain_text("high at noon")
    dialog.get_by_role("button", name="Ferries", exact=True).click()
    expect(dialog.get_by_role("row").filter(has_text="North pier")).to_contain_text("07:15")
    expect(dialog.get_by_text("Monday")).to_have_count(0)


def test_a_word_document_previews_as_its_page(page_for, make_user):
    page = page_for(make_user())
    letter = docx_bytes("Dear harbour master,", "The buoys need paint before May.")

    dialog = attach_and_open(page, "letter.docx", letter, DOCX)
    show_preview(dialog)

    page_view = dialog.locator("section.docx")
    expect(page_view).to_contain_text("The buoys need paint before May.")
    expect(page_view).to_contain_text("Dear harbour master,")


def titled_deck(*alignments: str | None) -> bytes:
    """One title slide per entry, its title set to that alignment or left to the layout."""
    pptx = pytest.importorskip("pptx")
    from pptx.enum.text import PP_ALIGN

    deck = pptx.Presentation()
    for alignment in alignments:
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = "Harbour"
        if alignment:
            slide.shapes.title.text_frame.paragraphs[0].alignment = getattr(PP_ALIGN, alignment)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def title_ink(dialog: Locator, number: int) -> tuple[float, float]:
    """Where the title's dark pixels start and end on a slide, as % of the slide width."""
    dialog.get_by_role("button", name=f"Slide {number}", exact=True).click()
    slide = dialog.get_by_role("img", name=f"Slide {number}", exact=True)
    expect(slide).to_be_visible()
    return tuple(
        slide.evaluate(
            """async image => {
                await image.decode();
                const canvas = document.createElement('canvas');
                canvas.width = image.naturalWidth;
                canvas.height = image.naturalHeight;
                const context = canvas.getContext('2d');
                context.drawImage(image, 0, 0);
                const { width, height } = canvas;
                const { data } = context.getImageData(0, 0, width, height);
                let first = width, last = -1;
                for (let y = Math.floor(height * 0.28); y < height * 0.48; y++) {
                    for (let x = 0; x < width; x++) {
                        const i = (y * width + x) * 4;
                        if (data[i] + data[i + 1] + data[i + 2] < 300) {
                            first = Math.min(first, x);
                            last = Math.max(last, x);
                        }
                    }
                }
                return [first / width * 100, last / width * 100];
            }"""
        )
    )


def test_a_slide_title_previews_with_the_alignment_the_file_gives_it(page_for, make_user):
    """#32873: a title the file aligns left was drawn centred, as if it had no alignment."""
    page = page_for(make_user())
    deck = titled_deck("LEFT", None, "RIGHT")

    dialog = attach_and_open(page, "harbour.pptx", deck, PPTX)
    show_preview(dialog)

    left_start, _ = title_ink(dialog, 1)
    assert left_start < 20, f"a left-aligned title starts {left_start:.0f}% across the slide"
    unset_start, unset_end = title_ink(dialog, 2)
    assert unset_start < unset_end and abs((unset_start + unset_end) / 2 - 50) < 5, (
        f"a title without alignment spans {unset_start:.0f}% to {unset_end:.0f}%, not centred"
    )
    _, right_end = title_ink(dialog, 3)
    assert right_end > 80, f"a right-aligned title ends {right_end:.0f}% across the slide"


def chart_deck(*chart_types: str) -> bytes:
    """One slide per chart type, each charting the same four shares."""
    pptx = pytest.importorskip("pptx")
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    data = CategoryChartData()
    data.categories = ["Cod", "Herring", "Mackerel", "Plaice"]
    data.add_series("Catch", (55, 18, 9, 18))
    deck = pptx.Presentation()
    for chart_type in chart_types:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_chart(
            getattr(XL_CHART_TYPE, chart_type), Inches(1), Inches(1), Inches(8), Inches(5.5), data
        )
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def chart_ink(dialog: Locator, number: int) -> dict[str, float]:
    """How much of a slide is in colour, and the widest single run of colour, as % of the slide."""
    dialog.get_by_role("button", name=f"Slide {number}", exact=True).click()
    slide = dialog.get_by_role("img", name=f"Slide {number}", exact=True)
    expect(slide).to_be_visible()
    return slide.evaluate(
        """async image => {
            await image.decode();
            const canvas = document.createElement('canvas');
            canvas.width = image.naturalWidth;
            canvas.height = image.naturalHeight;
            const context = canvas.getContext('2d');
            context.drawImage(image, 0, 0);
            const { width, height } = canvas;
            const { data } = context.getImageData(0, 0, width, height);
            let coloured = 0, widestRun = 0;
            for (let y = 0; y < height; y++) {
                let run = 0;
                for (let x = 0; x < width; x++) {
                    const i = (y * width + x) * 4;
                    const [r, g, b] = [data[i], data[i + 1], data[i + 2]];
                    const isColoured = Math.max(r, g, b) - Math.min(r, g, b) > 60;
                    coloured += isColoured ? 1 : 0;
                    run = isColoured ? run + 1 : 0;
                    widestRun = Math.max(widestRun, run);
                }
            }
            return {
                coloured: coloured / (width * height) * 100,
                widestRun: widestRun / width * 100
            };
        }"""
    )


def test_pie_doughnut_and_horizontal_bar_charts_preview_as_drawn_in_the_file(page_for, make_user):
    """#32884: pie and doughnut charts left an empty slide and horizontal bars stood upright."""
    page = page_for(make_user())
    deck = chart_deck("PIE", "DOUGHNUT", "BAR_CLUSTERED", "COLUMN_CLUSTERED")

    dialog = attach_and_open(page, "catch.pptx", deck, PPTX)
    show_preview(dialog)

    for number, kind in [(1, "pie"), (2, "doughnut")]:
        coloured = chart_ink(dialog, number)["coloured"]
        assert coloured > 10, f"the {kind} chart slide is only {coloured:.1f}% in colour"
    widest_bar = chart_ink(dialog, 3)["widestRun"]
    assert widest_bar > 40, f"the longest horizontal bar spans only {widest_bar:.0f}% of the slide"
    widest_column = chart_ink(dialog, 4)["widestRun"]
    assert widest_column < 20, f"a column of the column chart spans {widest_column:.0f}% across"


def test_markdown_previews_with_its_formatting(page_for, make_user):
    page = page_for(make_user())
    markdown = "# Tide table\n\n- **Monday**: high at noon\n- Tuesday: high at one\n"

    dialog = attach_and_open(page, "tides.md", markdown, "text/markdown")
    show_preview(dialog)

    expect(dialog.get_by_role("heading", name="Tide table")).to_be_visible()
    expect(dialog.get_by_role("listitem")).to_have_count(2)
    expect(dialog.locator("strong").filter(has_text="Monday")).to_be_visible()


def test_a_pdf_shows_the_text_read_out_of_it(page_for, make_user):
    page = page_for(make_user())

    dialog = attach_and_open(
        page, "chart.pdf", pdf_bytes("Soundings near the reef"), "application/pdf"
    )

    expect(dialog.get_by_text("1 Extracted Lines")).to_be_visible()
    expect(dialog.get_by_text("Soundings near the reef")).to_be_visible()
    expect(dialog.get_by_role("button", name="Preview", exact=True)).to_be_visible()


def red_buoy_png() -> bytes:
    picture = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 30, 30)).save(picture, format="PNG")
    return picture.getvalue()


def test_a_sent_picture_opens_full_size_and_downloads_as_uploaded(page_for, make_user, upstream):
    page = page_for(make_user())
    picture = red_buoy_png()
    attach(page, "buoy.png", picture, "image/png")
    expect(page.get_by_role("button", name="Show image preview")).to_be_visible()
    ask_with_attachments(page, upstream, f"what colour is the buoy? {uuid.uuid4().hex[:6]}")

    conversation(page).get_by_role("button", name="Show image preview").click()
    download_button = page.get_by_role("button", name="Download", exact=True)
    expect(download_button).to_be_visible()
    shown = page.locator(".modal img").last
    expect(shown).to_be_visible()
    assert shown.evaluate("image => [image.naturalWidth, image.naturalHeight]") == [40, 30]

    with page.expect_download() as downloaded:
        download_button.click()
    assert Path(downloaded.value.path()).read_bytes() == picture


def long_report() -> tuple[str, list[str]]:
    """Six paragraphs of about a chunk each, every one naming its own landmark."""
    landmarks = ["Hochosterwitz", "Landskron", "Griffen", "Taggenbrunn", "Glanegg", "Liebenfels"]
    paragraphs = [f"{name} castle. " + "The walls stand on the hill. " * 32 for name in landmarks]
    return "\n\n".join(paragraphs), landmarks


def landmarks_sent(upstream, question: str, landmarks: list[str]) -> list[str]:
    [request] = [body for body in upstream.chat_requests() if reply.answering(question)(body)]
    sent = json.dumps(request["messages"])
    return [name for name in landmarks if name in sent]


def test_a_files_mode_switch_shows_its_own_mode_after_another_file_is_removed(page_for, make_user):
    """#32699 (fix a944564c2): the remaining file's switch showed the removed file's mode."""
    page = page_for(make_user())
    report, _ = long_report()
    dialog = attach_and_open(page, "castles.txt", report, "text/plain")
    dialog.get_by_role("switch").click()
    expect(dialog.get_by_text("Using Entire Document")).to_be_visible()
    page.keyboard.press("Escape")
    expect(dialog).to_have_count(0)
    dialog = attach_and_open(page, "rivers.txt", "The Drau flows east.", "text/plain")
    expect(dialog.get_by_text("Using Focused Retrieval")).to_be_visible()
    page.keyboard.press("Escape")
    expect(dialog).to_have_count(0)

    removed = file_chip(page.locator("form"), "castles.txt")
    removed.hover()
    removed.get_by_role("button", name="Remove File").click()
    expect(page.locator("form").get_by_text("castles.txt")).to_have_count(0)

    dialog = open_file(page, file_chip(page.locator("form"), "rivers.txt"), "rivers.txt")
    expect(dialog.get_by_text("Using Focused Retrieval")).to_be_visible()
    expect(dialog.get_by_role("switch")).to_have_attribute("aria-checked", "false")


def test_using_the_entire_document_sends_the_model_all_of_it(page_for, make_user, upstream):
    page = page_for(make_user())
    report, landmarks = long_report()
    dialog = attach_and_open(page, "castles.txt", report, "text/plain")
    expect(dialog.get_by_text("Using Focused Retrieval")).to_be_visible()
    page.keyboard.press("Escape")
    focused_question = f"which castles? {uuid.uuid4().hex[:6]}"
    ask_with_attachments(page, upstream, focused_question)
    assert len(landmarks_sent(upstream, focused_question, landmarks)) < len(landmarks)

    page.goto("/")
    dialog = attach_and_open(page, "castles.txt", report, "text/plain")
    dialog.get_by_role("switch").click()
    expect(dialog.get_by_text("Using Entire Document")).to_be_visible()
    page.keyboard.press("Escape")
    whole_question = f"which castles, all of them? {uuid.uuid4().hex[:6]}"
    ask_with_attachments(page, upstream, whole_question)

    assert landmarks_sent(upstream, whole_question, landmarks) == landmarks
